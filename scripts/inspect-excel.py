from pathlib import Path
import openpyxl, collections, datetime, json, re, zipfile

ROOT=Path(__file__).resolve().parents[1]
source=ROOT/'data/ACRS_Data_Input.xlsm'
wv=openpyxl.load_workbook(source,read_only=True,data_only=True)
wf=openpyxl.load_workbook(source,read_only=False,data_only=False)
summary={}
for sv in wv:
    sf=wf[sv.title]
    rs=list(sv.values)
    forms=[(c.coordinate,c.value.text if hasattr(c.value,'text') else c.value) for row in sf for c in row if c.data_type=='f' or hasattr(c.value,'text')]
    errors=[c.coordinate+':'+str(c.value) for row in sv for c in row if c.data_type=='e']
    summary[sv.title]={'maxRows':sv.max_row,'maxCols':sv.max_column,'nonemptyRows':sum(any(v is not None for v in r) for r in rs),'formulas':len(forms),'formulaExamples':list(dict.fromkeys(v for _,v in forms))[:12],'errors':errors[:20],'errorCount':len(errors),'tables':{k:sf.tables[k].ref for k in sf.tables},'state':sf.sheet_state}
print(json.dumps(summary,ensure_ascii=False,indent=2,default=str))
points=[(i,r) for i,r in enumerate(wv['Ponto'].values,1) if i>1 and r[0] and r[1]]
bills=[(i,r) for i,r in enumerate(wv['Faturas'].values,1) if i>1 and r[1] and r[5] is not None]
print('PONTO',len(points),'FATURAS',len(bills))
print('PONTO date',min(r[1] for i,r in points),max(r[1] for i,r in points))
print('FAT date',min(r[0] for i,r in bills if isinstance(r[0],datetime.datetime)),max(r[0] for i,r in bills if isinstance(r[0],datetime.datetime)))
print('PONTO by month',collections.Counter(str(r[1])[:7] for i,r in points))
print('FAT by month',collections.Counter(str(r[0])[:7] for i,r in bills))
print('PONTO latest works',collections.Counter(str(r[8]) for i,r in points if str(r[1])[:7]=='2026-08'))
print('FAT latest works',collections.Counter(str(r[3]) for i,r in bills if str(r[0])[:7]=='2026-08'))
print('Ponto unique works',len({str(r[8]) for i,r in points}),'people',len({r[0] for i,r in points}))
print('Fatura unique works',len({str(r[3]) for i,r in bills}),'vendors',len({r[1] for i,r in bills}),'categories',len({r[4] for i,r in bills}),'total',sum(r[5] for i,r in bills if isinstance(r[5],(int,float))))
print('PONTO sample recent',points[-3:])
print('Named ranges',list(wf.defined_names))
print('Table refs',[(s.title,[(k,s.tables[k].ref) for k in s.tables]) for s in wf])
print('External links',len(wf._external_links))
with zipfile.ZipFile(source) as z:
    print('VBA',[n for n in z.namelist() if 'vba' in n.lower()])
    print('Connections',[n for n in z.namelist() if 'connection' in n.lower() or 'query' in n.lower()])
out=ROOT/'data/source';out.mkdir(exist_ok=True)
(out/'excel-structure.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2,default=str)+'\n')
